import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from openpyxl.utils import get_column_letter

def generate_soa_excel(data, scope_name="Périmètre ISO 27001", version_name="Draft"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Déclaration d'Applicabilité"

    ws.merge_cells('A1:F1')
    ws['A1'] = f"Déclaration d'Applicabilité (SoA) - {scope_name} (Version: {version_name})"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    headers = ["Code ISO", "Titre du Contrôle", "Thème", "Applicable", "Justification", "Statut"]
    ws.append([])
    ws.append(headers)

    header_fill = PatternFill(start_color='374151', end_color='374151', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    for col_num in range(1, 7):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    for entry in data:
        ctrl = entry.iso_control
        row = [
            getattr(ctrl, 'code', 'N/A'),
            getattr(ctrl, 'title', 'N/A'),
            getattr(ctrl, 'theme', 'N/A'),
            "Oui" if entry.is_applicable else "Non",
            entry.justification or "",
            entry.get_implementation_status_display() if hasattr(entry, 'get_implementation_status_display') else entry.implementation_status
        ]
        ws.append(row)
        
        current_row = ws.max_row
        for col_num in range(1, 7):
            c = ws.cell(row=current_row, column=col_num)
            c.border = thin_border
            if col_num in [1, 4]:
                c.alignment = Alignment(horizontal='center')

    # CORRECTION : Utilisation de get_column_letter + exclusion de la ligne de titre (A1) pour le calcul de taille
    for col_idx, col in enumerate(ws.columns, start=1):
        col_letter = get_column_letter(col_idx)
        
        # On calcule la longueur max en ignorant la première ligne (A1 fusionnée qui fausserait la largeur)
        max_len = max(len(str(cell.value or '')) for cell in col[1:])
        
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_soa_pdf(data, scope_name="Périmètre ISO 27001", version_name="Draft"):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )
    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1F2937'),
        spaceAfter=15
    )
    story.append(Paragraph(f"Déclaration d'Applicabilité (SoA) - {scope_name} | Version : <b>{version_name}</b>", title_style))

    headers = ["Code", "Titre du Contrôle", "Thème", "Applicable", "Justification", "Statut"]
    table_data = [[Paragraph(f"<b>{h}</b>", ParagraphStyle('HeaderStyle', parent=styles['Normal'], textColor=colors.white)) for h in headers]]

    cell_style = ParagraphStyle('TableCell', parent=styles['Normal'], fontSize=8, leading=10)

    for entry in data:
        ctrl = entry.iso_control
        row = [
            Paragraph(str(getattr(ctrl, 'code', 'N/A')), cell_style),
            Paragraph(str(getattr(ctrl, 'title', 'N/A')), cell_style),
            Paragraph(str(getattr(ctrl, 'theme', 'N/A')), cell_style),
            Paragraph("Oui" if entry.is_applicable else "Non", cell_style),
            Paragraph(entry.justification or "-", cell_style),
            Paragraph(entry.get_implementation_status_display() if hasattr(entry, 'get_implementation_status_display') else (entry.implementation_status or "-"), cell_style),
        ]
        table_data.append(row)

    col_widths = [50, 150, 100, 60, 250, 120]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#374151')),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,0), 8),
        ('TOPPADDING', (0,0), (-1,0), 8),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F9FAFB')]),
    ]))

    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()